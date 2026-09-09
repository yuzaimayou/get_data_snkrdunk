import os
import subprocess
import sys

import flet as ft
from openpyxl import load_workbook
from xlsx import snkr_check_file
from datetime import datetime

def main(page: ft.Page):
    page.title = "SNKRDUNK Checker"
    page.padding = 0
    page.bgcolor = ft.Colors.GREY_100
    page.window.width = 1100
    page.window.height = 720
    page.window.min_width = 900
    page.window.min_height = 600
    page.theme_mode = ft.ThemeMode.LIGHT

    card = dict(
        bgcolor=ft.Colors.WHITE,
        border_radius=14,
        padding=20,
        shadow=ft.BoxShadow(
            spread_radius=0,
            blur_radius=18,
            color=ft.Colors.with_opacity(0.08, ft.Colors.BLACK),
            offset=ft.Offset(0, 4),
        ),
    )

    # Header
    header = ft.Container(
        bgcolor=ft.Colors.WHITE,
        padding=ft.Padding.symmetric(horizontal=28, vertical=18),
        border=ft.Border.only(bottom=ft.BorderSide(1, ft.Colors.GREY_200)),
        content=ft.Row(
            alignment=ft.MainAxisAlignment.SPACE_BETWEEN,
            controls=[
                ft.Row(
                    spacing=12,
                    controls=[
                        ft.Container(
                            width=44,
                            height=44,
                            border_radius=12,
                            bgcolor=ft.Colors.INDIGO_50,
                            alignment=ft.Alignment.CENTER,
                            content=ft.Icon(ft.Icons.SHOPPING_BAG_OUTLINED, color=ft.Colors.INDIGO, size=24),
                        ),
                        ft.Column(
                            spacing=2,
                            controls=[
                                ft.Text("SNKRDUNK Checker", size=20, weight=ft.FontWeight.BOLD, color=ft.Colors.GREY_900),
                                ft.Text("Theo dõi và kiểm tra biến động giá", size=12, color=ft.Colors.GREY_600),
                            ],
                        ),
                    ]
                ),
                ft.Container(
                    bgcolor=ft.Colors.GREEN_50,
                    border_radius=20,
                    padding=ft.Padding.symmetric(horizontal=12, vertical=7),
                    content=ft.Row(
                        spacing=7,
                        controls=[
                            ft.Container(width=8, height=8, bgcolor=ft.Colors.GREEN, border_radius=10),
                            ft.Text("Sẵn sàng", size=12, weight=ft.FontWeight.W_600, color=ft.Colors.GREEN_800),
                        ],
                    ),
                ),
            ],
        ),
    )

    # Stats
    checked = ft.Text("0", size=25, weight=ft.FontWeight.BOLD)
    changed = ft.Text("0", size=25, weight=ft.FontWeight.BOLD)
    errors = ft.Text("0", size=25, weight=ft.FontWeight.BOLD)

    def stat_card(icon, title, value, subtitle):
        return ft.Container(
            expand=True,
            **card,
            content=ft.Row(
                spacing=14,
                controls=[
                    ft.Container(
                        width=42,
                        height=42,
                        border_radius=11,
                        bgcolor=ft.Colors.GREY_100,
                        alignment=ft.Alignment.CENTER,
                        content=ft.Icon(icon, size=21, color=ft.Colors.INDIGO_700),
                    ),
                    ft.Column(
                        spacing=2,
                        expand=True,
                        controls=[
                            ft.Text(title, size=12, color=ft.Colors.GREY_500),
                            value,
                            ft.Text(subtitle, size=10, color=ft.Colors.GREY_500),
                        ],
                    ),
                ],
            ),
        )

    stats = ft.Row(
        spacing=14,
        controls=[
            stat_card(ft.Icons.CHECK_CIRCLE_OUTLINE, "Đã kiểm tra", checked, "sản phẩm / size"),
            stat_card(ft.Icons.TRENDING_UP, "Có thay đổi", changed, "giá tăng hoặc giảm"),
            stat_card(ft.Icons.ERROR_OUTLINE, "Lỗi API", errors, "cần kiểm tra lại"),
        ],
    )

    # File picker
    selected_file = ft.Text("snkrdunk_results.xlsx", size=13, color=ft.Colors.GREY_800, max_lines=1)
    file_status = ft.Text("File đầu vào mặc định", size=11, color=ft.Colors.GREY_500)
    input_path = None
    output_path = None

    file_picker = ft.FilePicker()
    page.services.append(file_picker)

    async def pick_input_file(_):
        nonlocal input_path

        try:
            files = await file_picker.pick_files(
                dialog_title="Chọn file Excel",
                allow_multiple=False,
                file_type=ft.FilePickerFileType.CUSTOM,
                allowed_extensions=["xlsx", "xls"],
            )
            if files:
                input_path = files[0].path
                selected_file.value = os.path.basename(input_path) if input_path else files[0].name
                file_status.value = "Đã chọn file"
                file_status.color = ft.Colors.GREEN_700
            else:
                file_status.value = "Chưa chọn file"
                file_status.color = ft.Colors.GREY_500
            page.update()
        except Exception as exc:
            file_status.value = f"Không thể chọn file: {exc}"
            file_status.color = ft.Colors.RED_700
            page.update()

    progress = ft.ProgressBar(value=0, visible=False, height=5)
    status_text = ft.Text("Chưa bắt đầu kiểm tra", size=12, color=ft.Colors.GREY_600)

    def update_progress(current, total):
        if total > 0:
            progress.value = current / total
            status_text.value = f"Đang kiểm tra: {current}/{total} dòng"
        else:
            progress.value = 1
            status_text.value = "Không có dữ liệu để kiểm tra"
        page.update()

    def run_check():
        nonlocal input_path, output_path
        time=datetime.now().strftime("%d-%m-%y_%H-%M-%S")
        try:
            if not input_path:
                input_path = os.path.abspath("snkrdunk_results.xlsx")

            if not os.path.isfile(input_path):
                raise FileNotFoundError(f"Không tìm thấy file: {input_path}")

            output_path = os.path.join(
                os.path.dirname(input_path),
                f"data_modified_{time}.xlsx",
            )

            progress.visible = True
            progress.value = 0
            status_text.value = "Đang chuẩn bị kiểm tra..."
            start_button.disabled = True
            page.update()

            snkr_check_file(
                input_path,
                output_path,
                progress_callback=update_progress,
            )

            progress.value = 1
            status_text.value = f"Hoàn thành — đã lưu: {os.path.basename(output_path)}"
            result_wb = load_workbook(output_path, read_only=True)
            result_sheet = result_wb[result_wb.sheetnames[0]]
            checked.value = str(max(0, result_sheet.max_row - 1))
            result_wb.close()
            file_status.value = "Kiểm tra thành công"
            file_status.color = ft.Colors.GREEN_700

        except Exception as exc:
            progress.value = 0
            status_text.value = f"Kiểm tra thất bại: {exc}"
            file_status.value = "Có lỗi khi kiểm tra"
            file_status.color = ft.Colors.RED_700

        finally:
            start_button.disabled = False
            page.update()

    def start_check(_):
        page.run_thread(run_check)

    def reset_file(e):
        nonlocal input_path, output_path
        input_path = None
        output_path = None
        selected_file.value = "snkrdunk_results.xlsx"
        file_status.value = "File đầu vào mặc định"
        file_status.color = ft.Colors.GREY_500
        progress.visible = False
        progress.value = 0
        status_text.value = "Chưa bắt đầu kiểm tra"
        checked.value = "0"
        changed.value = "0"
        errors.value = "0"
        page.update()

    def open_result_file(_):
        if not output_path or not os.path.isfile(output_path):
            status_text.value = "Chưa có file kết quả để mở"
            page.update()
            return

        try:
            if sys.platform.startswith("win"):
                os.startfile(output_path)
            elif sys.platform == "darwin":
                subprocess.Popen(["open", output_path])
            else:
                subprocess.Popen(["xdg-open", output_path])
        except Exception as exc:
            status_text.value = f"Không thể mở file: {exc}"
            page.update()

    start_button = ft.FilledButton(
        "Bắt đầu kiểm tra",
        icon=ft.Icons.PLAY_ARROW_ROUNDED,
        height=46,
        style=ft.ButtonStyle(shape=ft.RoundedRectangleBorder(radius=10)),
        on_click=start_check,
    )

    file_card = ft.Container(
        **card,
        content=ft.Column(
            spacing=16,
            controls=[
                ft.Row(
                    alignment=ft.MainAxisAlignment.SPACE_BETWEEN,
                    controls=[
                        ft.Column(
                            spacing=3,
                            controls=[
                                ft.Text("File dữ liệu", size=16, weight=ft.FontWeight.BOLD, color=ft.Colors.GREY_900),
                                ft.Text("Chọn file Excel chứa danh sách sản phẩm cần kiểm tra.", size=12, color=ft.Colors.GREY_500),
                            ],
                        ),
                        ft.Icon(ft.Icons.TABLE_CHART_OUTLINED, color=ft.Colors.GREY_400),
                    ],
                ),
                ft.Container(
                    border=ft.Border.all(1, ft.Colors.GREY_200),
                    border_radius=10,
                    padding=ft.Padding.symmetric(horizontal=14, vertical=12),
                    content=ft.Row(
                        spacing=12,
                        controls=[
                            ft.Icon(ft.Icons.DESCRIPTION_OUTLINED, color=ft.Colors.INDIGO_600),
                            ft.Column(spacing=2, expand=True, controls=[selected_file, file_status]),
                            ft.TextButton(
                                "Chọn file",
                                icon=ft.Icons.UPLOAD_FILE,
                                on_click=pick_input_file,
                            ),
                            ft.IconButton(icon=ft.Icons.CLOSE, icon_size=18, tooltip="Đặt lại", on_click=reset_file),
                        ],
                    ),
                ),
                ft.Row(
                    alignment=ft.MainAxisAlignment.SPACE_BETWEEN,
                    controls=[ft.Row(spacing=10, controls=[status_text]), start_button],
                ),
                progress,
            ],
        ),
    )

    recent_file_card = ft.Container(
        **card,
        content=ft.Row(
            alignment=ft.MainAxisAlignment.SPACE_BETWEEN,
            controls=[
                ft.Row(
                    spacing=12,
                    controls=[
                        ft.Container(
                            width=42,
                            height=42,
                            border_radius=11,
                            bgcolor=ft.Colors.GREY_100,
                            alignment=ft.Alignment.CENTER,
                            content=ft.Icon(
                                ft.Icons.DESCRIPTION_OUTLINED,
                                size=21,
                                color=ft.Colors.INDIGO_700,
                            ),
                        ),
                        ft.Column(
                            spacing=3,
                            controls=[
                                ft.Text(
                                    "File kết quả",
                                    size=16,
                                    weight=ft.FontWeight.BOLD,
                                    color=ft.Colors.GREY_900,
                                ),
                                ft.Text(
                                    "Mở file kết quả gần nhất được tạo bởi chương trình.",
                                    size=12,
                                    color=ft.Colors.GREY_500,
                                ),
                            ],
                        ),
                    ],
                ),
                ft.TextButton(
                    "Mở file gần đây",
                    icon=ft.Icons.OPEN_IN_NEW,
                    on_click=open_result_file,
                ),
            ],
        ),
    )

    body = ft.Column(
        expand=True,
        spacing=18,
        scroll=ft.ScrollMode.AUTO,
        controls=[
            ft.Text("Tổng quan", size=24, weight=ft.FontWeight.BOLD, color=ft.Colors.GREY_900),
            ft.Text("Kiểm tra giá sản phẩm mới trên SNKRDUNK một cách nhanh chóng.", size=13, color=ft.Colors.GREY_600),
            stats,
            file_card,
            recent_file_card,
        ],
    )

    page.add(
        ft.Column(
            expand=True,
            spacing=0,
            controls=[
                header,
                ft.Container(expand=True, padding=ft.Padding.all(28), content=body),
            ],
        )
    )


if __name__ == "__main__":
    ft.run(main)
