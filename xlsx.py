from openpyxl import load_workbook
from database import get_or_fetch_url
from type import get_product_type
from request import get_data_products,get_data_apparels
def compare_price(data,size,price):
    try:
        
        if data and size!=None:
            size=str(size)
            if size not in data:
                return "Không tìm thấy size"
            if data[size]==None and price!=None:
                return "Hàng đã bị bán"
            elif price==None:
                return ""
            elif data[size]<price:
                return "Giá đã giảm"

            elif data[size]>price:
                return "Giá đã tăng"
        return ""
    except Exception as e:
        print(data, size)
        print(e)
        return ""
    
def snkr_check_file(input_path,output_path, progress_callback=None):
    wb=load_workbook(input_path)
    sheet=wb[wb.sheetnames[0]]
    maxRow=sheet.max_row

    data_list={}

    for i in range(2,maxRow+1):
        id=sheet.cell(row=i,column=1).value
        name=sheet.cell(row=i,column=2).value
        size=sheet.cell(row=i,column=3).value
        price=sheet.cell(row=i,column=4).value

        product_key=(id,name)

        if product_key not in data_list:
            product_type=get_product_type(id)

            if(product_type=="products"):
                data_list[product_key]=get_data_products(id)
            else:
                data_list[product_key]=get_data_apparels(get_or_fetch_url(id,name))

        note=compare_price(data_list[product_key],size,price)

        sheet.cell(row=i,column=5,value=note)
        if progress_callback:
            progress_callback(i-1,maxRow-1)


    wb.save(output_path)
