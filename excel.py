from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from models import SizeListing


HEADERS = ["Scanned At", "Size", "Min Listing Price", "Listing Item Count"]


class ExcelStorageError(Exception):
    """Raised when listing data cannot be saved to an Excel workbook."""


def _sheet_name(product_id: str) -> str:
    """Return an Excel-safe worksheet name for a product ID."""
    # API validation already excludes Excel-invalid sheet-name characters.
    # Excel imposes a maximum worksheet-name length of 31 characters.
    return product_id[:31]


def _style_header(ws) -> None:
    """Apply simple formatting to the worksheet header."""
    fill = PatternFill(fill_type="solid", fgColor="D9E2F3")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill


def append_listings_to_excel(
    file_path: str,
    product_id: str,
    listings: list[SizeListing],
) -> int:
    """Append one scan to the worksheet belonging to ``product_id``.

    A missing worksheet is created with headers in row 1 and listing data from
    row 2. If it already exists, new listing rows are appended immediately after
    the last populated row. Each size is one Excel row, so one scan can add
    multiple consecutive rows.
    """
    path = Path(file_path).expanduser()
    if path.suffix.lower() != ".xlsx":
        raise ExcelStorageError("Please select an .xlsx Excel file.")
    if not path.exists():
        raise ExcelStorageError("The selected Excel file does not exist.")

    workbook = None
    try:
        workbook = load_workbook(path)
        sheet_name = _sheet_name(product_id)

        if sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            if ws.max_row == 1 and all(ws.cell(1, col).value is None for col in range(1, 5)):
                for col, header in enumerate(HEADERS, start=1):
                    ws.cell(1, col, header)
                _style_header(ws)
        else:
            ws = workbook.create_sheet(title=sheet_name)
            for col, header in enumerate(HEADERS, start=1):
                ws.cell(1, col, header)
            _style_header(ws)

        next_row = max(ws.max_row + 1, 2)
        scanned_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        for index, item in enumerate(listings):
            row = next_row + index
            ws.cell(row, 1, scanned_at)
            ws.cell(row, 2, item.size)
            ws.cell(row, 3, item.min_listing_price)
            ws.cell(row, 4, item.listing_item_count)

        ws.column_dimensions[get_column_letter(1)].width = 20
        ws.column_dimensions[get_column_letter(2)].width = 18
        ws.column_dimensions[get_column_letter(3)].width = 22
        ws.column_dimensions[get_column_letter(4)].width = 22
        ws.freeze_panes = "A2"

        workbook.save(path)
        return len(listings)
    except PermissionError as exc:
        raise ExcelStorageError(
            "Cannot save the Excel file. Please close it if it is open in another program."
        ) from exc
    except OSError as exc:
        raise ExcelStorageError(f"Unable to save the Excel file: {exc}") from exc
    except Exception as exc:
        raise ExcelStorageError(f"Unable to update the Excel file: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()
